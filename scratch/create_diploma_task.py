import os, sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls

doc = docx.Document()

# Standard margins (A4: top 2.0cm, bottom 2.0cm, left 2.5cm, right 1.5cm)
for section in doc.sections:
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.79)
    section.bottom_margin = Inches(0.79)
    section.left_margin = Inches(0.98)
    section.right_margin = Inches(0.59)

normal_style = doc.styles['Normal']
normal_style.font.name = 'Times New Roman'
normal_style.font.size = Pt(13)
normal_style.font.color.rgb = RGBColor(0, 0, 0)
normal_style.paragraph_format.line_spacing = 1.15
normal_style.paragraph_format.space_after = Pt(3)
normal_style.paragraph_format.space_before = Pt(0)

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_table_borders(table, color="B0B0B0"):
    tblPr = table._tbl.tblPr
    borders_xml = f"""
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
        <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
        <w:left w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
        <w:right w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
        <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
        <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
    </w:tblBorders>
    """
    tblPr.append(parse_xml(borders_xml))

def add_p(text="", align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, italic=False, space_before=0, space_after=3, indent=False, size=13):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.49)
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
    return p

# ШАПКА
add_p("МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12)
add_p("ДЕРЖАВНИЙ ВИЩИЙ НАВЧАЛЬНИЙ ЗАКЛАД", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12)
add_p("«УЖГОРОДСЬКИЙ НАЦІОНАЛЬНИЙ УНІВЕРСИТЕТ»", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13)
add_p("ФАКУЛЬТЕТ ІНФОРМАЦІЙНИХ ТЕХНОЛОГІЙ", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12)
add_p("КАФЕДРА ІНФОРМАЦІЙНИХ УПРАВЛЯЮЧИХ СИСТЕМ ТА ТЕХНОЛОГІЙ", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, space_after=10)

# ГРИФ ЗАТВЕРДЖЕННЯ
p_appr = add_p(align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=14)
p_appr.add_run("ЗАТВЕРДЖУЮ\n").bold = True
p_appr.add_run("Завідувач кафедри ІУСТ\nд.т.н., професор Олександр МІЦА\n")
p_appr.add_run("«___» ________________ 2026 року")

# ЗАГОЛОВОК
add_p("ЗАВДАННЯ", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=15)
add_p("НА ДИПЛОМНУ РОБОТУ МАГІСТРА", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14)
add_p("Стефаніву Олегу Тарасовичу", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, italic=True, size=14, space_after=12)

# ПУНКТИ ЗАВДАННЯ
p1 = add_p()
p1.add_run("1. Тема проекту (роботи): ").bold = True
p1.add_run("«Інтелектуальна система оцінки кардіоваскулярного ризику на основі методів пояснюваного штучного інтелекту».")

p2 = add_p()
p2.add_run("2. Строк здачі студентом закінченого проекту (роботи): ").bold = True
p2.add_run("20 грудня 2026 року.")

p3 = add_p()
p3.add_run("3. Вихідні дані до проекту (роботи): ").bold = True
p3.add_run("Клінічні набори даних кардіологічних спостережень (CDC BRFSS, CardioTrain N = 68 414), настанови Європейського товариства кардіологів (ESC SCORE2, ESC 2024), алгоритми машинного навчання (XGBoost, EBM, Random Forest) та фреймворки інтерпретованості (TreeSHAP).")

p4 = add_p()
p4.add_run("4. Зміст пояснювальної записки (перелік питань, які підлягають розробці):").bold = True

questions = [
    "Опрацювати наукову літературу, проаналізувати існуючі клінічні шкали стратифікації ризику серцево-судинних ускладнень (SCORE2, Framingham) та комп'ютерні системи підтримки прийняття рішень (CDSS).",
    "Дослідити та адаптувати алгоритми машинного навчання (градієнтний бустинг XGBoost, Explainable Boosting Machine, випадковий ліс, логістичну регресію) для задачі класифікації кардіоваскулярного ризику.",
    "Здійснити порівняльний бенчмаркінг методів пояснюваного штучного інтелекту (TreeSHAP, LIME, KernelSHAP) за показниками вірності (Deletion-AUC), обчислювальної латентності та рангової стійкості.",
    "Розробити алгоритм інтерактивного контрфактичного моделювання («What-If») для генерації персоніфікованих цільових рекомендацій з оптимізації ризику.",
    "Спроектувати архітектуру та програмно реалізувати багаторівневу веб-платформу CardioRisk AI з модулем генерації консультативних медичних звітів у форматі PDF відповідно до протоколів ESC."
]
for q in questions:
    pq = add_p(f"– {q}", indent=True, size=12)

p5 = add_p(space_before=4)
p5.add_run("5. Перелік графічного матеріалу (з точним зазначенням обов'язкових ілюстрацій):").bold = True
gfx = [
    "Архітектурна діаграма дворівневої системи клінічної підтримки рішень CardioRisk AI.",
    "Дерево клінічного аналізу та маршрутизації пацієнта (скринінг → лабораторне дообстеження).",
    "Графіки порівняльного аналізу моделей машинного навчання (ROC- та PR-криві).",
    "Водоспадні діаграми факторного аналізу TreeSHAP Waterfall та карти важливості біомаркерів.",
    "Розробка мультимедійної презентації для представлення та захисту магістерської роботи (15–20 слайдів)."
]
for g in gfx:
    pg = add_p(f"– {g}", indent=True, size=12)

# КАЛЕНДАРНИЙ ПЛАН
add_p("6. Календарний план", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13, space_before=8, space_after=6)

table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(table)

col_widths = [Inches(0.6), Inches(3.8), Inches(1.6), Inches(1.1)]

headers = ["№ п/п", "Назва етапів дипломного проекту (роботи)", "Строк виконання етапів проекту", "Примітки"]
hdr_cells = table.rows[0].cells
for i, h in enumerate(headers):
    hdr_cells[i].text = h
    hdr_cells[i].width = col_widths[i]
    set_cell_shading(hdr_cells[i], "E2E8F0")
    p = hdr_cells[i].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.bold = True
    p.runs[0].font.size = Pt(11)
    p.runs[0].font.name = 'Times New Roman'

plan_data = [
    ("1.", "Розробка і затвердження детального плану та індивідуального завдання дипломної роботи", "до 20 вересня 2026 р.", "Виконано"),
    ("2.", "Збір, структурування та аналіз наукової літератури, клінічних протоколів ESC та медичних наборів даних", "до 30 вересня 2026 р.", "Виконано"),
    ("3.", "Написання теоретичного розділу: огляд існуючих клінічних шкал, моделей ML та методів Explainable AI", "до 15 жовтня 2026 р.", "Виконано"),
    ("4.", "Розробка, навчання та валідація моделей (XGBoost, EBM), науковий бенчмаркінг методів XAI (TreeSHAP vs LIME)", "до 30 жовтня 2026 р.", "У процесі"),
    ("5.", "Архітектурне проектування та програмна реалізація веб-платформи CardioRisk AI з What-If симулятором та генератором PDF", "до 15 листопада 2026 р.", "У процесі"),
    ("6.", "Проведення експериментальної оцінки системи, оформлення пояснювальної записки та підготовка до передзахисту", "до 30 листопада 2026 р.", "Заплановано")
]

for row_data in plan_data:
    row_cells = table.add_row().cells
    for i, val in enumerate(row_data):
        row_cells[i].text = val
        row_cells[i].width = col_widths[i]
        p = row_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in [0, 2, 3] else WD_ALIGN_PARAGRAPH.LEFT
        p.runs[0].font.size = Pt(11)
        p.runs[0].font.name = 'Times New Roman'

add_p(space_after=6)
p_date = add_p()
p_date.add_run("Дата видачі завдання: ").bold = True
p_date.add_run("17 вересня 2026 року")

# ПІДПИСИ
add_p(space_after=8)
p_signs = doc.add_table(rows=1, cols=2)
p_signs.alignment = WD_TABLE_ALIGNMENT.CENTER
cell_l, cell_r = p_signs.rows[0].cells

p_l = cell_l.paragraphs[0]
p_l.add_run("Студент: _________________ Олег СТЕФАНІВ\n")
p_l.runs[0].font.size = Pt(12)

p_r = cell_r.paragraphs[0]
p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p_r.add_run("Науковий керівник: _________________ Володимир НІКОЛЕНКО\n")
p_r.runs[0].font.size = Pt(12)

out_task_path = r"C:\Users\Олег\Desktop\Магістерська\Документи\Завдання_на_дипломну_роботу_Стефанів.docx"
doc.save(out_task_path)
print(f"Diploma task saved successfully to: {out_task_path} ({os.path.getsize(out_task_path)} bytes)")
