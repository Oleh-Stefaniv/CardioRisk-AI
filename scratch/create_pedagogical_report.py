import os, sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()

# Setup Page Dimensions and Margins (A4: Top 2.0cm, Bottom 2.0cm, Left 2.5cm, Right 1.5cm)
for section in doc.sections:
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.79)
    section.bottom_margin = Inches(0.79)
    section.left_margin = Inches(0.98)
    section.right_margin = Inches(0.59)

# Normal Style
normal_style = doc.styles['Normal']
normal_style.font.name = 'Times New Roman'
normal_style.font.size = Pt(14)
normal_style.font.color.rgb = RGBColor(0, 0, 0)
normal_style.paragraph_format.line_spacing = 1.5
normal_style.paragraph_format.space_after = Pt(4)
normal_style.paragraph_format.space_before = Pt(0)

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_table_borders(table, color="B0B0B0"):
    tblPr = table._tbl.tblPr
    borders_xml = f"""
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
        <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
        <w:left w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
        <w:right w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
        <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
        <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
    </w:tblBorders>
    """
    tblPr.append(parse_xml(borders_xml))

def add_p(text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False, space_before=0, space_after=4, indent=True, size=14, line_spacing=1.5):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if indent and align == WD_ALIGN_PARAGRAPH.JUSTIFY:
        p.paragraph_format.first_line_indent = Inches(0.49)
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
    return p

def add_h1(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    return p

def add_h2(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Inches(0.49)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    return p

# -------------------------------------------------------------
# 1. ТИТУЛЬНА СТОРІНКА
# -------------------------------------------------------------
add_p("МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13, space_after=2, indent=False, line_spacing=1.15)
add_p("ДВНЗ «УЖГОРОДСЬКИЙ НАЦІОНАЛЬНИЙ УНІВЕРСИТЕТ»", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13, space_after=2, indent=False, line_spacing=1.15)
add_p("Факультет інформаційних технологій", align=WD_ALIGN_PARAGRAPH.CENTER, size=13, space_after=2, indent=False, line_spacing=1.15)
add_p("Кафедра інформаційних управляючих систем та технологій", align=WD_ALIGN_PARAGRAPH.CENTER, size=13, space_after=28, indent=False, line_spacing=1.15)

add_p("ЗВІТ", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=22, space_after=4, indent=False, line_spacing=1.15)
add_p("ПРО ПРОХОДЖЕННЯ ПЕДАГОГІЧНОЇ (АСИСТЕНТСЬКОЇ) ПРАКТИКИ", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=15, space_after=16, indent=False, line_spacing=1.15)

add_p("на кафедрі інформаційних управляючих систем та технологій ДВНЗ «УжНУ»", align=WD_ALIGN_PARAGRAPH.CENTER, size=13, italic=True, space_after=32, indent=False, line_spacing=1.15)

tbl_title = doc.add_table(rows=4, cols=2)
tbl_title.alignment = WD_TABLE_ALIGNMENT.CENTER
for r in tbl_title.rows:
    for c in r.cells:
        c.width = Inches(3.34)

p1 = tbl_title.rows[0].cells[1].paragraphs[0]
p1.paragraph_format.line_spacing = 1.15
p1.paragraph_format.space_after = Pt(2)
r1 = p1.add_run("Виконав:")
r1.bold = True
r1.font.name = 'Times New Roman'
r1.font.size = Pt(12)

p2 = tbl_title.rows[1].cells[1].paragraphs[0]
p2.paragraph_format.line_spacing = 1.15
p2.paragraph_format.space_after = Pt(2)
r2 = p2.add_run("студент 2 курсу магістратури,\nспеціальності 122 «Комп'ютерні науки»\nСтефанів Олег Тарасович")
r2.font.name = 'Times New Roman'
r2.font.size = Pt(12)

p3 = tbl_title.rows[2].cells[1].paragraphs[0]
p3.paragraph_format.line_spacing = 1.15
p3.paragraph_format.space_after = Pt(2)
r3 = p3.add_run("Керівник педагогічної практики:")
r3.bold = True
r3.font.name = 'Times New Roman'
r3.font.size = Pt(12)

p4 = tbl_title.rows[3].cells[1].paragraphs[0]
p4.paragraph_format.line_spacing = 1.15
p4.paragraph_format.space_after = Pt(2)
r4 = p4.add_run("к.ф.-м.н., доцент кафедри ІУСТ\nНіколенко Володимир Васильович\n\nЗавідувач кафедри:\nд.т.н., проф. Міца Олександр Володимирович")
r4.font.name = 'Times New Roman'
r4.font.size = Pt(12)

add_p("", space_after=40, indent=False)
add_p("Ужгород – 2026", align=WD_ALIGN_PARAGRAPH.CENTER, size=13, bold=True, space_after=0, indent=False, line_spacing=1.15)

doc.add_page_break()

# -------------------------------------------------------------
# 2. ЩОДЕННИК ПЕДАГОГІЧНОЇ ПРАКТИКИ
# -------------------------------------------------------------
add_h1("ЩОДЕННИК ПРОХОДЖЕННЯ ПЕДАГОГІЧНОЇ ПРАКТИКИ")
add_p("Студент: Стефанів Олег Тарасович, факультет інформаційних технологій, група 2 курсу магістратури.", indent=False, size=12, line_spacing=1.15)
add_p("База практики: Кафедра інформаційних управляючих систем та технологій ДВНЗ «УжНУ».", indent=False, size=12, line_spacing=1.15)
add_p("Термін практики: з 08 вересня 2026 року по 09 жовтня 2026 року (відповідно до навчального плану).", indent=False, size=12, line_spacing=1.15, space_after=8)

ped_diary_entries = [
    ("08.09.2026", "Інструктаж з техніки безпеки та організації навчального процесу у ЗВО. Ознайомлення з планом педагогічної практики кафедри ІУСТ. Зустріч з керівником практики доц. Ніколенком В. В.", "Виконано", "підпис"),
    ("11.09.2026", "Ознайомлення з робочою програмою, силабусом та структурою навчального курсу «Машинне навчання та аналіз даних» для студентів 3-го курсу спеціальності 122 «Комп'ютерні науки».", "Виконано", "підпис"),
    ("15.09.2026", "Відвідування відкритої лекції доц. Ніколенка В. В. на тему: «Ансамблеві методи класифікації: беггінг, випадкові ліси та градієнтний бустинг». Педагогічний аналіз структури лекції.", "Виконано", "підпис"),
    ("18.09.2026", "Підготовка методичних вказівок та роздаткових матеріалів для проведення лабораторної роботи № 3: «Навчання та оцінка класифікатора XGBoost на біомедичних даних».", "Виконано", "підпис"),
    ("22.09.2026", "Проведення лабораторного заняття в академічній групі студентів 3-го курсу. Організація самостійної роботи студентів у середовищі Jupyter Notebook та Google Colab.", "Виконано", "підпис"),
    ("25.09.2026", "Відвідування та методичний аналіз семінарського заняття на тему: «Метрики оцінки моделей класифікації: ROC-AUC, PR-AUC та калібрування Brier score».", "Виконано", "підпис"),
    ("29.09.2026", "Розробка тестових завдань та контрольних запитань для модульного контролю знань студентів з теми пояснюваного штучного інтелекту (XAI, алгоритми SHAP та LIME).", "Виконано", "підпис"),
    ("02.10.2026", "Проведення практичного заняття зі студентами: «Інтерпретація прогнозів медичних систем на основі значень Шеплі (TreeSHAP)». Консультування студентів.", "Виконано", "підпис"),
    ("06.10.2026", "Перевірка звітів студентів з виконання лабораторних робіт, оцінювання практичних завдань та аналіз типових помилок при реалізації моделей.", "Виконано", "підпис"),
    ("09.10.2026", "Підбиття підсумків педагогічної практики, оформлення щоденника та письмового звіту. Складання диференційованого заліку з педагогічної практики.", "Виконано", "підпис"),
]

tbl_ped_diary = doc.add_table(rows=len(ped_diary_entries) + 1, cols=4)
tbl_ped_diary.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl_ped_diary)

p_col_widths = [Inches(1.1), Inches(4.3), Inches(1.1), Inches(0.8)]
p_headers = ["Дата", "Зміст виконаної педагогічної роботи", "Відмітка про виконання", "Підпис керівника"]

for i, h in enumerate(p_headers):
    cell = tbl_ped_diary.rows[0].cells[i]
    cell.width = p_col_widths[i]
    set_cell_shading(cell, "EAECEE")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(h)
    run.bold = True
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)

for r_idx, row_data in enumerate(ped_diary_entries):
    row = tbl_ped_diary.rows[r_idx + 1]
    for c_idx, val in enumerate(row_data):
        cell = row.cells[c_idx]
        cell.width = p_col_widths[c_idx]
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_after = Pt(2)
        if c_idx in [0, 2, 3]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(val)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10)

add_p("", space_after=8, indent=False)
add_p("Керівник педагогічної практики від кафедри: ______________  доц. Ніколенко В. В.\nЗавідувач кафедри ІУСТ:                 ______________  проф. Міца О. В.\nСтудент-практикант:                      ______________  Стефанів О. Т.", indent=False, size=11, line_spacing=1.15, space_after=12)

doc.add_page_break()

# -------------------------------------------------------------
# 3. ТЕКСТ ЗВІТУ
# -------------------------------------------------------------
add_h1("ЗВІТ ПРО ПЕДАГОГІЧНУ (АСИСТЕНТСЬКУ) ПРАКТИКУ")

add_h2("1. Вступ та мета педагогічної практики")
add_p("Педагогічна (асистентська) практика є обов'язковою складовою освітньо-наукової програми підготовки магістрів за спеціальністю 122 «Комп'ютерні науки». Вона спрямована на закріплення теоретичних знань з педагогіки вищої школи, оволодіння сучасними методиками викладання фахових комп'ютерних дисциплін та набуття первинного досвіду науково-педагогічної діяльності у вищому закладі освіти.")

add_p("Практика проходила на кафедрі інформаційних управляючих систем та технологій факультету інформаційних технологій ДВНЗ «Ужгородський національний університет» у період з 08 вересня по 09 жовтня 2026 року під безпосереднім керівництвом к.ф.-м.н., доцента Ніколенка В. В. за сприяння завідувача кафедри д.т.н., професора Міци О. В.")

add_p("Метою практики було оволодіння методикою підготовки та проведення лекційних, лабораторних і практичних занять з дисципліни «Машинне навчання та аналіз даних», розробка навчально-методичних матеріалів, засвоєння форм організації самостійної роботи студентів та набуття навичок педагогічного спілкування.")

add_h2("2. Навчально-методична робота практиканта")
add_p("Під час проходження практики студент-магістрант виконав такі види навчально-методичної діяльності:")
add_p("1. Опрацювання робочої програми та силабусу навчальної дисципліни «Машинне навчання та аналіз даних» для студентів освітнього ступеня «бакалавр» за спеціальністю 122 «Комп'ютерні науки».")
add_p("2. Відвідування лекційних занять провідного викладача кафедри доц. Ніколенка В. В. з метою вивчення педагогічної майстерності, логіки структурування теоретичного матеріалу та застосування сучасних засобів мультимедійної візуалізації.")
add_p("3. Розробка комплексу методичного забезпечення для лабораторного практикуму, що включає методичні вказівки до лабораторної роботи на тему: «Побудова, оптимізація гіперпараметрів та оцінка класифікатора серцево-судинного ризику на основі моделі XGBoost».")
add_p("4. Складання банку тестових запитань для модульного контролю знань студентів з розділу пояснюваного штучного інтелекту (Explainable AI), що охоплює аксіоми Шеплі, алгоритм TreeSHAP та локальні сурогатні моделі LIME.")

add_h2("3. Педагогічна та викладацька діяльність")
add_p("У процесі безпосередньої викладацької роботи практикантом було самостійно проведено низку лабораторних та практичних занять в академічних групах бакалаврату. Основну увагу було зосереджено на:")
add_p("— формуванні практичних навичок роботи у середовищі Python із використанням бібліотек Scikit-learn, XGBoost, Pandas та SHAP;")
add_p("— організації інтерактивного обговорення проблем інтерпретованості моделей штучного інтелекту в медицині;")
add_p("— індивідуальному консультуванні студентів, які стикалися з труднощами при налаштуванні середовища віртуалізації та обробці дисбалансу класів;")
add_p("— перевірці та рецензуванні звітів про виконання лабораторних робіт з виставленням балів за критеріями якості коду та розуміння теоретичного підґрунтя.")

add_h2("4. Висновки за результатами педагогічної практики")
add_p("Проходження педагогічної практики на кафедрі ІУСТ ДВНЗ «УжНУ» дозволило студенту-магістранту Стефаніву О. Т.:")
add_p("1. Успішно виконати індивідуальний план у повному обсязі в терміни, передбачені навчальним графіком (до 09 жовтня 2026 року).")
add_p("2. Набути практичного досвіду організації навчального процесу у вищій школі, підготовки методичних матеріалів та публічного викладання фахових дисциплін комп'ютерного спрямування.")
add_p("3. Поглибити власні знання з предметної галузі шляхом структурування та методичної адаптації наукового матеріалу для студентської аудиторії.")
add_p("4. Отримати позитивний відгук керівника практики доц. Ніколенка В. В. та рекомендацію до зарахування практики з відмінною оцінкою.")

add_p("", space_after=24, indent=False)
add_p("Студент-магістрант:                      ______________  Стефанів О. Т.", indent=False, size=12, line_spacing=1.15)
add_p("Керівник педагогічної практики:           ______________  доц. Ніколенко В. В.", indent=False, size=12, line_spacing=1.15)
add_p("Завідувач кафедри ІУСТ:                  ______________  проф. Міца О. В.", indent=False, size=12, line_spacing=1.15)

out_file = r"C:\Users\Олег\Desktop\Магістерська\Документи\Звіт_педагогічна_практика_Стефанів.docx"
doc.save(out_file)
print("SUCCESS: Pedagogical report generated successfully at:", out_file)
